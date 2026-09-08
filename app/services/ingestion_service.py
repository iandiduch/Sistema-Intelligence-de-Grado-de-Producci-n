"""Ingesta de documentos: creacion del job (persistencia + encolado) y,
para el worker, parsing async por formato + chunking. Ninguno de los formatos
tiene una libreria async nativa, asi que cada parser sincronico corre en un
thread aparte via asyncio.to_thread para no bloquear el event loop.

La creacion del job vive aca (no en el endpoint) para que app/api/v1/endpoints
mantenga los handlers delgados: HTTP/multipart es lo unico que le corresponde
a la capa API, el resto -- persistir el registro y encolarlo -- es logica de
servicio."""

import asyncio
import logging
import re
import uuid
from pathlib import Path
from uuid import UUID

import pdfplumber
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from redis.asyncio import Redis
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import Settings
from app.core.exceptions import DocumentParsingError, UnsupportedFileTypeError
from app.db.models_orm import IngestionJob
from app.domain.models import FileType, TaskStatus
from app.schemas.ingest import ChunkMetadata, ChunkWithMetadata, PageContent

logger = logging.getLogger(__name__)


async def create_ingestion_job(
    db: AsyncSession,
    redis: Redis,
    settings: Settings,
    *,
    job_id: UUID,
    filename: str,
    file_type: FileType,
    storage_path: Path,
) -> IngestionJob:
    """Persiste el job en PENDING y lo encola. Si el encolado falla, el job
    ya quedo persistido -- se marca FAILED en la misma operacion en vez de
    dejar un job huerfano que nunca va a ser tomado por el worker."""
    job = IngestionJob(
        job_id=job_id,
        filename=filename,
        file_type=file_type.value,
        storage_path=str(storage_path),
        status=TaskStatus.PENDING.value,
    )
    db.add(job)
    await db.commit()
    await db.refresh(job)

    try:
        await redis.rpush(settings.REDIS_INGEST_QUEUE_KEY, str(job_id))
    except Exception as exc:  # noqa: BLE001 - Si falla el broker Redis, se actualiza el estado del job a FAILED
        job.status = TaskStatus.FAILED.value
        job.error_message = f"No se pudo encolar el job de ingesta: {exc}"
        await db.commit()
        logger.error("ingestion_service.enqueue_failed", extra={"job_id": str(job_id), "error": str(exc)})

    return job


async def list_ingestion_jobs(
    db: AsyncSession,
    *,
    status: TaskStatus | None = None,
    file_type: FileType | None = None,
    limit: int = 50,
    offset: int = 0,
) -> list[IngestionJob]:
    stmt = select(IngestionJob).order_by(IngestionJob.created_at.desc()).limit(limit).offset(offset)
    if status is not None:
        stmt = stmt.where(IngestionJob.status == status.value)
    if file_type is not None:
        stmt = stmt.where(IngestionJob.file_type == file_type.value)

    result = await db.execute(stmt)
    return list(result.scalars().all())


_HEADING_RE = re.compile(r"^#{1,6}\s+(.*)$", re.MULTILINE)


async def parse_document(storage_path: Path, file_type: FileType) -> list[PageContent]:
    parser = {
        FileType.PDF: _parse_pdf,
        FileType.DOCX: _parse_docx,
        FileType.TXT: _parse_txt,
        FileType.MD: _parse_md,
    }.get(file_type)
    if parser is None:
        raise UnsupportedFileTypeError(f"Formato no soportado: {file_type}")

    try:
        return await asyncio.to_thread(parser, storage_path)
    except (OSError, ValueError) as exc:
        raise DocumentParsingError(f"No se pudo parsear {storage_path.name}: {exc}") from exc


def _parse_pdf(path: Path) -> list[PageContent]:
    pages: list[PageContent] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(PageContent(page=i, section=None, text=text))
    return pages


def _parse_docx(path: Path) -> list[PageContent]:
    doc = DocxDocument(str(path))
    pages: list[PageContent] = []
    current_section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            pages.append(PageContent(page=None, section=current_section, text="\n".join(buffer)))
            buffer.clear()

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style is not None and paragraph.style.name.startswith("Heading"):
            flush()
            current_section = text
            continue
        buffer.append(text)
    flush()

    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        table_text = "\n".join(r for r in rows if r.strip())
        if table_text:
            pages.append(PageContent(page=None, section=current_section, text=table_text))

    return pages


def _parse_txt(path: Path) -> list[PageContent]:
    text = path.read_text(encoding="utf-8")
    return [PageContent(page=None, section=None, text=text)] if text.strip() else []


def _parse_md(path: Path) -> list[PageContent]:
    raw = path.read_text(encoding="utf-8")
    headings = list(_HEADING_RE.finditer(raw))
    if not headings:
        return [PageContent(page=None, section=None, text=raw)] if raw.strip() else []

    pages: list[PageContent] = []
    for i, match in enumerate(headings):
        start = match.end()
        end = headings[i + 1].start() if i + 1 < len(headings) else len(raw)
        body = raw[start:end].strip()
        if body:
            pages.append(PageContent(page=None, section=match.group(1).strip(), text=body))
    return pages


def build_chunks(
    pages: list[PageContent],
    document_id: uuid.UUID,
    filename: str,
    file_type: FileType,
    settings: Settings,
) -> list[ChunkWithMetadata]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        separators=["\n## ", "\n### ", "\n\n", "\n", ". ", " ", ""],
    )

    chunks: list[ChunkWithMetadata] = []
    counter = 0
    for page in pages:
        for piece in splitter.split_text(page.text):
            if not piece.strip():
                continue
            chunks.append(
                ChunkWithMetadata(
                    chunk_id=f"{document_id}_{counter}",
                    text=piece,
                    metadata=ChunkMetadata(
                        document_id=str(document_id),
                        filename=filename,
                        file_type=file_type,
                        page=page.page,
                        section=page.section,
                        source=filename,
                    ),
                )
            )
            counter += 1
    return chunks
